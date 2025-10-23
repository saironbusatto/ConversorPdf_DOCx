import os
import sys
import webbrowser
from flask import Flask, render_template, request, send_file, jsonify
import pdfplumber
from docx import Document
from pathlib import Path
import threading

app = Flask(__name__)

# Detectar se está rodando como executável
if getattr(sys, 'frozen', False):
    BASE_DIR = Path(sys.executable).parent
else:
    BASE_DIR = Path(__file__).parent

UPLOAD_FOLDER = BASE_DIR / 'uploads'
OUTPUT_FOLDER = BASE_DIR / 'Documentos convertidos'

UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)



@app.route('/')
def index():
    return render_template('index.html')

@app.route('/convert', methods=['POST'])
def convert():
    if 'files[]' not in request.files:
        return jsonify({'error': 'Nenhum arquivo enviado'}), 400
    
    files = request.files.getlist('files[]')
    converted_files = []
    
    for file in files:
        if file.filename == '':
            continue
            
        if not file.filename.lower().endswith('.pdf'):
            continue
        
        pdf_path = UPLOAD_FOLDER / file.filename
        file.save(pdf_path)
        
        docx_filename = file.filename.rsplit('.', 1)[0] + '.docx'
        docx_path = OUTPUT_FOLDER / docx_filename
        
        try:
            doc = Document()
            
            with pdfplumber.open(str(pdf_path)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text and text.strip():
                        doc.add_paragraph(text)
            
            doc.save(str(docx_path))
            converted_files.append(docx_filename)
        except Exception as e:
            print(f"Erro ao converter {file.filename}: {e}")
        finally:
            try:
                pdf_path.unlink()
            except:
                pass
    
    return jsonify({'files': converted_files})

@app.route('/view/<filename>')
def view(filename):
    import re
    file_path = OUTPUT_FOLDER / filename
    doc = Document(file_path)
    text = '\n\n'.join([para.text for para in doc.paragraphs if para.text.strip()])
    
    # Encontrar todas as ocorrências de PROCESSO:
    matches = list(re.finditer(r'PROCESSO:', text, re.IGNORECASE))
    tags = [{'index': i, 'position': m.start()} for i, m in enumerate(matches)]
    
    return jsonify({'text': text, 'tags': tags})

@app.route('/download/<filename>')
def download(filename):
    file_path = OUTPUT_FOLDER / filename
    return send_file(file_path, as_attachment=True)

def open_browser():
    webbrowser.open('http://127.0.0.1:5000')

if __name__ == '__main__':
    threading.Timer(1, open_browser).start()
    app.run(debug=False, port=5000)
